"""
web/exports.py -- build the Excel (.xlsx) export.

Sheets, in order:
  Records      -- fixed date/item/qty/price/total rows (OCR fusion, mobile app)
  All pages    -- per document with more than one page: every page's rows
                  in one filterable table with a Page column, one block per
                  distinct set of headings
  Page 1, 2... -- one sheet per page, that page's tables stacked

With *tidy*, numbers written with thousands separators ("15 000", "5,000")
become real numbers and day-first dates ("29.11.17", "20|12|17") become real
dates, so they sort, filter and add up in Excel. Without it, cells stay
exactly as written except plain numbers like "2.50".
"""

from __future__ import annotations

import io
import re
from datetime import datetime
from pathlib import Path

_PLAIN_NUMBER = re.compile(r"-?(0|[1-9]\d*)(\.\d+)?")
_GROUPED_NUMBER = re.compile(r"-?\d{1,3}(?:[ ,]\d{3})+(\.\d+)?")
_DATE = re.compile(r"(\d{1,2})\s*[./|\-]\s*(\d{1,2})\s*[./|\-]\s*(\d{2}|\d{4})")
_SHEET_BAD_CHARS = re.compile(r"[\[\]:*?/\\]")
_TIMESTAMP_PREFIX = re.compile(r"^\d{8}T\d+_")


def document_label(scan: dict) -> str:
    """Human name for a saved document: its given name, else the file name."""
    if scan.get("name"):
        return scan["name"]
    return _TIMESTAMP_PREFIX.sub("", Path(scan.get("source_file") or f"scan {scan['id']}").stem)


def _parse_date(text: str) -> datetime | None:
    m = _DATE.fullmatch(text.strip())
    if not m:
        return None
    day, month, year = (int(g) for g in m.groups())
    if len(m.group(3)) == 2:
        year += 2000 if year <= (datetime.now().year % 100) + 1 else 1900
    try:
        return datetime(year, month, day)
    except ValueError:
        return None


def _date_columns(rows: list[list[str]]) -> set[int]:
    """Columns where most filled cells are day-first dates."""
    cols = set()
    for c in range(max((len(r) for r in rows), default=0)):
        cells = [r[c] for r in rows if c < len(r) and r[c].strip()]
        if cells and sum(_parse_date(x) is not None for x in cells) / len(cells) >= 0.6:
            cols.add(c)
    return cols


def _convert(text: str, tidy: bool, is_date_col: bool):
    """Value to write for one cell (number, date or the text itself)."""
    value = text.strip()
    if _PLAIN_NUMBER.fullmatch(value):
        return float(value) if "." in value else int(value)
    if tidy:
        if _GROUPED_NUMBER.fullmatch(value):
            digits = re.sub(r"[ ,]", "", value)
            return float(digits) if "." in digits else int(digits)
        if is_date_col:
            parsed = _parse_date(value)
            if parsed:
                return parsed
    return text


def build_workbook(records: list[dict], record_columns: list[str], scans: list[dict], tidy: bool = True) -> bytes:
    from openpyxl import Workbook
    from openpyxl.styles import Font, PatternFill

    title_font = Font(bold=True, size=12)
    header_font = Font(bold=True, color="FFFFFF")
    header_fill = PatternFill("solid", fgColor="0B1040")

    wb = Workbook()
    wb.remove(wb.active)
    used_names: set[str] = set()

    def new_sheet(name: str):
        name = _SHEET_BAD_CHARS.sub("-", name).strip() or "Sheet"
        base, n = name[:31], 2
        while name[:31].lower() in used_names:
            suffix = f" ({n})"
            name = base[: 31 - len(suffix)] + suffix
            n += 1
        used_names.add(name[:31].lower())
        return wb.create_sheet(name[:31])

    def write_header(ws, values):
        ws.append(list(values))
        for cell in ws[ws.max_row]:
            cell.font, cell.fill = header_font, header_fill
        return ws.max_row

    def write_title(ws, text):
        ws.append([text])
        ws[ws.max_row][0].font = title_font

    def write_rows(ws, rows, lead: list | None = None):
        """Rows of text cells; *lead* holds per-row values prepended as-is (e.g. page)."""
        date_cols = _date_columns(rows) if tidy else set()
        for i, row in enumerate(rows):
            values = [_convert(v, tidy, c in date_cols) for c, v in enumerate(row)]
            prefix = [lead[i]] if lead else []
            ws.append(prefix + values)
            for cell, value, raw in zip(ws[ws.max_row][len(prefix):], values, row):
                if isinstance(value, float):
                    decimals = len(raw.strip().split(".")[1]) if "." in raw else 0
                    cell.number_format = "#,##0." + "0" * decimals if decimals else "#,##0"
                elif isinstance(value, int) and _GROUPED_NUMBER.fullmatch(raw.strip()):
                    cell.number_format = "#,##0"
                elif isinstance(value, datetime):
                    cell.number_format = "DD/MM/YYYY"

    def fit_columns(ws):
        widths: dict[str, int] = {}
        for row in ws.iter_rows():
            for cell in row:
                if cell.value is not None:
                    text = cell.value.strftime("%d/%m/%Y") if isinstance(cell.value, datetime) else str(cell.value)
                    widths[cell.column_letter] = max(widths.get(cell.column_letter, 0), len(text))
        for letter, width in widths.items():
            ws.column_dimensions[letter].width = min(width + 2, 50)

    if records or not scans:
        ws = new_sheet("Records")
        write_header(ws, record_columns)
        write_rows(ws, [[str(r[c]) if r[c] is not None else "" for c in record_columns] for r in records])
        ws.freeze_panes = "A2"
        if records:
            ws.auto_filter.ref = ws.dimensions
        fit_columns(ws)

    for scan in scans:
        label = document_label(scan)
        single = len(scans) == 1
        pages: dict[int, list[dict]] = {}
        for table in scan["tables"]:
            pages.setdefault(table.get("page") or 1, []).append(table)

        if len(pages) > 1:
            ws = new_sheet("All pages" if single else f"{label[:23]} all")
            groups: dict[tuple, list] = {}
            for page in sorted(pages):
                for table in pages[page]:
                    key = tuple(c.strip().lower() for c in table["columns"])
                    groups.setdefault(key, [table["columns"], [], []])
                    groups[key][1].extend(table["rows"])
                    groups[key][2].extend([page] * len(table["rows"]))
            for i, (columns, rows, page_numbers) in enumerate(groups.values()):
                if i:
                    ws.append([])
                header_row = write_header(ws, ["Page"] + list(columns))
                write_rows(ws, rows, lead=page_numbers)
                if i == 0:
                    ws.freeze_panes = f"A{header_row + 1}"
                    if len(groups) == 1 and rows:
                        ws.auto_filter.ref = f"A{header_row}:{ws.cell(row=ws.max_row, column=len(columns) + 1).coordinate}"
            fit_columns(ws)

        for page in sorted(pages):
            ws = new_sheet(f"Page {page}" if single else f"{label[:24]} p{page}")
            for i, table in enumerate(pages[page]):
                if i:
                    ws.append([])
                if table["title"]:
                    write_title(ws, table["title"])
                header_row = write_header(ws, table["columns"])
                if i == 0:
                    ws.freeze_panes = f"A{header_row + 1}"
                write_rows(ws, table["rows"])
            fit_columns(ws)

    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()
