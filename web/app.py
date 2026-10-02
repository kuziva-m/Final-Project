"""
web/app.py -- Flask front end for the OCR ledger pipeline, serving both the
HTML review UI and a JSON API for the companion mobile app.

No new OCR/preprocessing/extraction logic lives here; this module only wires
together the existing pipeline:

    preprocessing.clean.preprocess -- deskew/denoise/CLAHE/binarise
    ocr.fusion.run_engines         -- Tesseract-on-cleaned + EasyOCR-on-raw
    ocr.fusion.fuse_from_texts     -- fused into structured rows + provenance

See ocr/fusion.py for the fusion rule: Tesseract's rows are the structural
skeleton (date, item, column layout); EasyOCR supplies clean numeric values
where the two disagree, and every decision is logged (provenance) so the
result is auditable -- no model, only rules.

When ANTHROPIC_API_KEY is set, ocr/claude_reader.py reads the photo instead
and returns rows in the same shape; if that call fails, the fusion above is
used. The engine that produced the rows is always reported as used_engine.

Run:
    python -m web.app
    (or) python web/app.py
"""

from __future__ import annotations

import csv
import io
import json
import os
import re
import sqlite3
import sys
from datetime import datetime, timezone
from pathlib import Path

# Allow `python web/app.py` to find the top-level packages (preprocessing,
# ocr, extraction) regardless of the process's working directory.
_REPO_ROOT = Path(__file__).resolve().parent.parent
if str(_REPO_ROOT) not in sys.path:
    sys.path.insert(0, str(_REPO_ROOT))

import cv2
from flask import (
    Flask,
    Response,
    flash,
    g,
    jsonify,
    redirect,
    render_template,
    request,
    send_from_directory,
    url_for,
)
from werkzeug.exceptions import RequestEntityTooLarge
from werkzeug.utils import secure_filename

from extraction.fields import FIELDNAMES
from ocr import claude_reader
from ocr.fusion import fuse_from_texts, run_engines
from preprocessing.clean import preprocess
from preprocessing.pdf_pages import PdfError, render_pages

# ---------------------------------------------------------------------------
# Paths -- everything reads from / writes to folders, nothing is hardcoded.
# ---------------------------------------------------------------------------

DATA_DIR = _REPO_ROOT / "data"
UPLOADS_DIR = DATA_DIR / "uploads"
PROCESSED_DIR = DATA_DIR / "processed"
DB_PATH = DATA_DIR / "records.db"

for _dir in (UPLOADS_DIR, PROCESSED_DIR):
    _dir.mkdir(parents=True, exist_ok=True)

ALLOWED_EXTENSIONS = {".png", ".jpg", ".jpeg", ".tif", ".tiff", ".bmp", ".webp", ".pdf"}

app = Flask(__name__)
# Scanned PDFs can be large; reject anything over the limit with a clear message.
app.config["MAX_CONTENT_LENGTH"] = int(os.environ.get("LEDGER_MAX_UPLOAD_MB", "100")) * 1024 * 1024
# Flask's 500 KB default for submitted form fields is too small for a review
# page holding many pages of tables.
app.config["MAX_FORM_MEMORY_SIZE"] = 16 * 1024 * 1024
app.secret_key = os.environ.get("FLASK_SECRET_KEY", "dev-secret-key-change-in-production")


# ---------------------------------------------------------------------------
# SQLite -- stdlib sqlite3, table auto-created on startup.
# ---------------------------------------------------------------------------


def init_db() -> None:
    """Create the tables if they do not already exist.

    records -- fixed date/item/qty/price/total rows (OCR fusion, mobile app).
    scans   -- documents read with their own headings; tables_json holds
               [{"title", "columns", "rows"}], since columns vary per business.
    """
    with sqlite3.connect(str(DB_PATH)) as conn:
        conn.execute(
            """
            CREATE TABLE IF NOT EXISTS records (
                id          INTEGER PRIMARY KEY AUTOINCREMENT,
                date        TEXT,
                item        TEXT,
                qty         TEXT,
                price       TEXT,
                total       TEXT,
                source_file TEXT,
                created_at  TEXT NOT NULL
            )
            """
        )
        conn.execute(
            """
            CREATE TABLE IF NOT EXISTS scans (
                id          INTEGER PRIMARY KEY AUTOINCREMENT,
                source_file TEXT,
                engine      TEXT,
                tables_json TEXT NOT NULL,
                created_at  TEXT NOT NULL
            )
            """
        )
        conn.commit()


def get_db() -> sqlite3.Connection:
    """Return the request-scoped SQLite connection, creating it if needed."""
    if "db" not in g:
        g.db = sqlite3.connect(str(DB_PATH))
        g.db.row_factory = sqlite3.Row
    return g.db


@app.teardown_appcontext
def close_db(_exc: BaseException | None) -> None:
    db = g.pop("db", None)
    if db is not None:
        db.close()


# ---------------------------------------------------------------------------
# Helpers
# ---------------------------------------------------------------------------


def _allowed_file(filename: str) -> bool:
    return Path(filename).suffix.lower() in ALLOWED_EXTENSIONS


def _unique_stem(original_filename: str) -> str:
    """Timestamp-prefixed, path-safe stem so repeat uploads never collide."""
    safe_name = secure_filename(original_filename)
    stem = Path(safe_name).stem or "upload"
    timestamp = datetime.now(timezone.utc).strftime("%Y%m%dT%H%M%S%f")
    return f"{timestamp}_{stem}"


def _env_int(name: str, default: int) -> int:
    try:
        return max(1, int(os.environ.get(name, default)))
    except ValueError:
        return default


# A scanned record book can be dozens of pages; each page read by Claude
# costs a few cents, so cap the pages read per PDF. The free OCR fallback
# takes far longer per page on CPU, so it reads fewer.
MAX_PDF_PAGES = _env_int("LEDGER_MAX_PDF_PAGES", 40)
MAX_FALLBACK_PDF_PAGES = _env_int("LEDGER_MAX_FALLBACK_PDF_PAGES", 5)


def _pages_note(read: int, total: int, is_pdf: bool) -> str:
    if not is_pdf:
        return ""
    if read < total:
        return f"first {read} of {total} pages"
    return f"{total} page" + ("s" if total != 1 else "")


def _save_upload(upload) -> tuple[Path | None, str | None]:
    """Validate and save *upload*; return ``(saved_path, None)`` or ``(None, error)``."""
    if upload is None or upload.filename == "":
        return None, "No file selected. Choose a photo or PDF and try again."

    if not _allowed_file(upload.filename):
        return None, f"Unsupported file type. Allowed: {', '.join(sorted(ALLOWED_EXTENSIONS))}"

    stem = _unique_stem(upload.filename)
    ext = Path(secure_filename(upload.filename)).suffix.lower()
    source_path = UPLOADS_DIR / f"{stem}{ext}"
    upload.save(str(source_path))
    return source_path, None


def _uploaded_path(token: str) -> Path | None:
    """The saved upload named by *token* (from /upload), or None if invalid."""
    name = secure_filename(token or "")
    if not name or name != token or not _allowed_file(name):
        return None
    path = UPLOADS_DIR / name
    return path if path.is_file() else None


def _run_pipeline_on_upload(upload, detect_tables: bool = False) -> tuple[dict | None, str | None]:
    """Save *upload*, then run the pipeline on it (see _run_pipeline_on_file)."""
    source_path, error = _save_upload(upload)
    if error:
        return None, error
    return _run_pipeline_on_file(source_path, detect_tables)


def _run_pipeline_on_file(source_path: Path, detect_tables: bool = False) -> tuple[dict | None, str | None]:
    """Run preprocess + recognition on a saved upload and return the result.

    Shared by the HTML /process route and the JSON /api/scan route so both
    surfaces run the exact same pipeline. Returns ``(artifacts, None)`` on
    success or ``(None, error_message)`` on a validation/decode failure.

    A PDF is rendered to one image per page first; every page then goes
    through the same per-image pipeline, and the first page is the preview.

    ``artifacts`` keys: used_engine, source_file, raw_filename,
    processed_filename, pages_note, plus either ``tables`` (when
    *detect_tables* and Claude read the pages with their own headings) or
    ``rows`` + ``provenance`` (fixed FIELDNAMES rows).
    """
    source_file = source_path.name
    stem = source_path.stem
    ext = source_path.suffix.lower()

    is_pdf = ext == ".pdf"
    if is_pdf:
        try:
            page_paths, total_pages = render_pages(source_path, UPLOADS_DIR, stem, MAX_PDF_PAGES)
        except PdfError as exc:
            source_path.unlink(missing_ok=True)
            return None, str(exc)
    else:
        page_paths, total_pages = [source_path], 1

    def clean_page(page_path: Path) -> Path | None:
        img = cv2.imread(str(page_path), cv2.IMREAD_UNCHANGED)
        if img is None:
            return None
        clean_path = PROCESSED_DIR / f"{page_path.stem}_clean.png"
        preprocess(img, save_path=clean_path)
        return clean_path

    first_clean = clean_page(page_paths[0])
    if first_clean is None:
        source_path.unlink(missing_ok=True)
        return None, "Could not read that image. It may be corrupt or an unsupported format."

    files = {
        "source_file": source_file,
        "raw_filename": page_paths[0].name,
        "processed_filename": first_clean.name,
    }

    rows = None
    fallback_note = ""
    if claude_reader.is_enabled():
        claude_engine = f"Claude vision ({claude_reader.model_name()})"
        try:
            if detect_tables:
                tables = claude_reader.read_tables_from_pages(page_paths)
                if not tables:
                    tables = [{"title": "", "columns": ["Column 1", "Column 2", "Column 3"], "rows": [["", "", ""]]}]
                return {**files, "tables": tables, "used_engine": claude_engine,
                        "pages_note": _pages_note(len(page_paths), total_pages, is_pdf)}, None
            rows = claude_reader.read_rows_from_pages(page_paths)
            provenance = [{} for _ in rows]
            used_engine = claude_engine
            pages_read = len(page_paths)
        except claude_reader.ClaudeReadError as exc:
            app.logger.warning("Claude read failed, falling back to OCR fusion: %s", exc)
            fallback_note = f" -- Claude unavailable: {exc}"

    if rows is None:
        # Tesseract-on-cleaned (row structure) + EasyOCR-on-raw (clean values),
        # each engine called once per page, then fused with provenance -- which
        # engine won each field -- so the review UI can show what fusion
        # auto-corrected.
        rows, provenance = [], []
        fallback_pages = page_paths[:MAX_FALLBACK_PDF_PAGES]
        for i, page_path in enumerate(fallback_pages):
            clean_path = first_clean if i == 0 else clean_page(page_path)
            if clean_path is None:
                continue
            tesseract_text, easyocr_text = run_engines(page_path, processed_path=clean_path)
            page_rows, page_prov = fuse_from_texts(tesseract_text, easyocr_text, return_provenance=True)
            rows.extend(page_rows)
            provenance.extend(page_prov)
        used_engine = "fused (tesseract + easyocr)" + fallback_note
        pages_read = len(fallback_pages)

    if not rows:
        # Never dead-end the owner with nothing to correct -- give one blank row.
        rows = [{k: "" for k in FIELDNAMES}]
        provenance = [{}]
        used_engine = "none (manual entry)"

    return {**files, "rows": rows, "provenance": provenance, "used_engine": used_engine,
            "pages_note": _pages_note(pages_read, total_pages, is_pdf)}, None


def _save_rows(source_file: str, rows: list[dict]) -> int:
    """Insert non-blank *rows* into SQLite under *source_file*; return count saved."""
    created_at = datetime.now(timezone.utc).isoformat()
    db = get_db()
    saved = 0
    for row in rows:
        clean = {field: str(row.get(field, "") or "").strip() for field in FIELDNAMES}
        if not any(clean.values()):
            continue  # skip fully-blank rows (e.g. an unused manual-entry row)
        db.execute(
            """
            INSERT INTO records (date, item, qty, price, total, source_file, created_at)
            VALUES (:date, :item, :qty, :price, :total, :source_file, :created_at)
            """,
            {**clean, "source_file": source_file, "created_at": created_at},
        )
        saved += 1
    db.commit()
    return saved


@app.errorhandler(RequestEntityTooLarge)
def too_large(_exc):
    limit_mb = app.config["MAX_CONTENT_LENGTH"] // (1024 * 1024)
    message = f"That file is too large. The limit is {limit_mb} MB."
    if request.path.startswith("/api/") or request.path == "/upload":
        return jsonify({"error": message}), 413
    flash(message)
    return redirect(url_for("index"))


# ---------------------------------------------------------------------------
# HTML routes
# ---------------------------------------------------------------------------


@app.route("/")
def index():
    """Upload form for a single ledger image."""
    return render_template("index.html")


def _render_review(artifacts: dict):
    return render_template(
        "process.html",
        tables=artifacts.get("tables"),
        rows=artifacts.get("rows"),
        source_file=artifacts["source_file"],
        raw_filename=artifacts["raw_filename"],
        processed_filename=artifacts["processed_filename"],
        used_engine=artifacts["used_engine"],
        pages_note=artifacts["pages_note"],
    )


def _review_path(token: str) -> Path | None:
    """Where the read result for upload *token* is kept, or None if invalid."""
    name = secure_filename(token or "")
    if not name or name != token:
        return None
    return PROCESSED_DIR / f"{name}.review.json"


@app.route("/process", methods=["POST"])
def process():
    """Run the pipeline and show the editable table(s).

    Plain form submit (no JavaScript): takes the file and renders the review.

    With a ``token`` from /upload (the page's script uploads first so it can
    show real upload progress): reads that file, keeps the result, and
    answers ``{"review_url": ...}`` -- or ``{"error": ...}`` -- so the script
    can open the review page as a normal page load.
    """
    token = request.form.get("token")
    if token:
        source_path = _uploaded_path(token)
        if source_path is None:
            return jsonify({"error": "That upload could not be found. Please choose the file again."}), 404
        artifacts, error = _run_pipeline_on_file(source_path, detect_tables=True)
        if error:
            return jsonify({"error": error}), 400
        _review_path(token).write_text(json.dumps(artifacts, ensure_ascii=False), encoding="utf-8")
        return jsonify({"review_url": url_for("review", token=token)})

    artifacts, error = _run_pipeline_on_upload(request.files.get("image"), detect_tables=True)
    if error:
        flash(error)
        return redirect(url_for("index"))
    return _render_review(artifacts)


@app.route("/review/<token>")
def review(token: str):
    """Show a result read earlier via /process (refresh-safe)."""
    path = _review_path(token)
    if path is None or not path.is_file():
        flash("That review could not be found. Please upload the file again.")
        return redirect(url_for("index"))
    return _render_review(json.loads(path.read_text(encoding="utf-8")))


@app.route("/upload", methods=["POST"])
def upload():
    """Save an upload without reading it; returns ``{"token": ...}`` for /process."""
    source_path, error = _save_upload(request.files.get("image"))
    if error:
        return jsonify({"error": error}), 400
    return jsonify({"token": source_path.name})


@app.route("/save", methods=["POST"])
def save():
    """Persist the (possibly corrected) rows from /process into SQLite."""
    source_file = request.form.get("source_file", "").strip()
    try:
        num_rows = int(request.form.get("num_rows", "0"))
    except ValueError:
        num_rows = 0

    rows = [
        {field: request.form.get(f"{field}_{i}", "") for field in FIELDNAMES}
        for i in range(num_rows)
    ]
    saved = _save_rows(source_file, rows)

    flash(f"Saved {saved} record(s) from {source_file or 'upload'}.")
    return redirect(url_for("index"))


# Bounds on the dynamic review form, so a crafted POST can't make the server
# loop over millions of fields.
_MAX_TABLES, _MAX_COLS, _MAX_ROWS, _MAX_PAGE = 200, 50, 2000, 10000


def _form_int(name: str, upper: int) -> int:
    try:
        return max(0, min(int(request.form.get(name, "0")), upper))
    except ValueError:
        return 0


@app.route("/save-tables", methods=["POST"])
def save_tables():
    """Persist tables read with the page's own headings into the scans table."""
    source_file = request.form.get("source_file", "").strip()
    engine = request.form.get("engine", "").strip()

    tables = []
    for t in range(_form_int("num_tables", _MAX_TABLES)):
        n_cols = _form_int(f"num_cols_{t}", _MAX_COLS)
        n_rows = _form_int(f"num_rows_{t}", _MAX_ROWS)
        columns = [request.form.get(f"col_{t}_{c}", "").strip() for c in range(n_cols)]
        rows = []
        for r in range(n_rows):
            row = [request.form.get(f"cell_{t}_{r}_{c}", "").strip() for c in range(n_cols)]
            if any(row):
                rows.append(row)
        if rows:
            tables.append({
                "title": request.form.get(f"title_{t}", "").strip(),
                "columns": [col or f"Column {i + 1}" for i, col in enumerate(columns)],
                "rows": rows,
                "page": _form_int(f"page_{t}", _MAX_PAGE) or 1,
            })

    if not tables:
        flash("Nothing to save: every row was empty.")
        return redirect(url_for("index"))

    db = get_db()
    db.execute(
        "INSERT INTO scans (source_file, engine, tables_json, created_at) VALUES (?, ?, ?, ?)",
        (source_file, engine, json.dumps(tables, ensure_ascii=False), datetime.now(timezone.utc).isoformat()),
    )
    db.commit()

    n_rows = sum(len(t["rows"]) for t in tables)
    flash(f"Saved {n_rows} row(s) in {len(tables)} table(s) from {source_file or 'upload'}.")
    return redirect(url_for("index"))


# ---------------------------------------------------------------------------
# JSON API -- used by the companion mobile app (app/). Same pipeline and
# SQLite store as the HTML routes above, just JSON in/out instead of
# server-rendered forms.
# ---------------------------------------------------------------------------


@app.route("/api/health")
def api_health():
    """Liveness probe -- the mobile app pings this to detect a cold Space waking up."""
    return jsonify({"status": "ok"})


@app.route("/api/scan", methods=["POST"])
def api_scan():
    """Upload a ledger photo, run the pipeline, return rows + provenance as JSON."""
    artifacts, error = _run_pipeline_on_upload(request.files.get("image"))
    if error:
        return jsonify({"error": error}), 400

    return jsonify({
        "rows": artifacts["rows"],
        "provenance": artifacts["provenance"],
        "used_engine": artifacts["used_engine"],
        "pages_note": artifacts["pages_note"],
        "source_file": artifacts["source_file"],
        "raw_url": url_for("uploaded_file", filename=artifacts["raw_filename"]),
        "processed_url": url_for("processed_file", filename=artifacts["processed_filename"]),
    })


@app.route("/api/save", methods=["POST"])
def api_save():
    """Persist (possibly corrected) rows into SQLite. Body: {source_file, rows}."""
    body = request.get_json(silent=True) or {}
    source_file = str(body.get("source_file", "")).strip()
    rows = body.get("rows", [])
    if not isinstance(rows, list):
        return jsonify({"error": "'rows' must be a list."}), 400

    saved = _save_rows(source_file, rows)
    return jsonify({"saved": saved})


_NUMBER_RE = re.compile(r"-?(0|[1-9]\d*)(\.\d+)?")
_SHEET_BAD_CHARS = re.compile(r"[\[\]:*?/\\]")
_TIMESTAMP_PREFIX = re.compile(r"^\d{8}T\d+_")


def _cell_value(text: str):
    """Plain numbers become real numbers (so sums work in Excel); anything
    else -- dates like 12/03, codes like 007, words -- stays text."""
    if _NUMBER_RE.fullmatch(text):
        return float(text) if "." in text else int(text)
    return text


def _excel_workbook(records: list[dict], record_columns: list[str], scans: list[dict]) -> bytes:
    """Build the .xlsx export: one sheet per page of each scan, plus a
    Records sheet for the fixed-field rows. Tables from the same page are
    stacked on that page's sheet, each under its own title and header row."""
    from openpyxl import Workbook
    from openpyxl.styles import Font, PatternFill

    title_font = Font(bold=True, size=12)
    header_font = Font(bold=True, color="FFFFFF")
    header_fill = PatternFill("solid", fgColor="0B1040")

    wb = Workbook()
    wb.remove(wb.active)
    used_names: set[str] = set()

    def new_sheet(name: str):
        name = _SHEET_BAD_CHARS.sub("-", name).strip() or "Sheet"
        base, n = name[:31], 2
        while name[:31].lower() in used_names:
            suffix = f" ({n})"
            name = base[: 31 - len(suffix)] + suffix
            n += 1
        used_names.add(name[:31].lower())
        return wb.create_sheet(name[:31])

    def write_row(ws, values, font=None, fill=None, numbers=False):
        ws.append([_cell_value(v) if numbers else v for v in values])
        for cell in ws[ws.max_row]:
            if font:
                cell.font = font
            if fill:
                cell.fill = fill
            if isinstance(cell.value, float):
                decimals = len(str(values[cell.column - 1]).split(".")[1])
                cell.number_format = "0." + "0" * decimals

    def fit_columns(ws):
        widths: dict[str, int] = {}
        for row in ws.iter_rows():
            for cell in row:
                if cell.value is not None:
                    widths[cell.column_letter] = max(widths.get(cell.column_letter, 0), len(str(cell.value)))
        for letter, width in widths.items():
            ws.column_dimensions[letter].width = min(width + 2, 50)

    if records or not scans:
        ws = new_sheet("Records")
        write_row(ws, record_columns, header_font, header_fill)
        for record in records:
            write_row(ws, [str(record[c] if record[c] is not None else "") for c in record_columns], numbers=True)
        fit_columns(ws)

    for scan in scans:
        label = _TIMESTAMP_PREFIX.sub("", Path(scan["source_file"] or f"scan {scan['id']}").stem)
        pages: dict[int, list[dict]] = {}
        for table in scan["tables"]:
            pages.setdefault(table.get("page") or 1, []).append(table)
        for page in sorted(pages):
            ws = new_sheet(f"Page {page}" if len(scans) == 1 else f"{label[:24]} p{page}")
            for i, table in enumerate(pages[page]):
                if i:
                    ws.append([])
                if table["title"]:
                    write_row(ws, [table["title"]], title_font)
                write_row(ws, table["columns"], header_font, header_fill)
                for row in table["rows"]:
                    write_row(ws, row, numbers=True)
            fit_columns(ws)

    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


@app.route("/export/<fmt>")
def export(fmt: str):
    """Download everything saved as Excel, CSV or JSON.

    Covers both the fixed-field records and the scans saved with their own
    headings. Excel puts each page of each scan on its own sheet (see
    _excel_workbook). CSV lists the records first, then each scanned table as
    its own block (a label line, its header row, its rows). JSON is
    {"records": [...], "scans": [...]}.

    The download is named after the source image when everything exported
    came from the same upload; a flash confirms exactly what was exported
    (visible on the next page view, since a file download doesn't navigate
    the browser away from the current page).
    """
    if fmt not in ("xlsx", "csv", "json"):
        flash(f"Unknown export format '{fmt}'. Use 'xlsx', 'csv' or 'json'.")
        return redirect(url_for("index"))

    db = get_db()
    columns = ["id", "date", "item", "qty", "price", "total", "source_file", "created_at"]
    records = [dict(row) for row in db.execute(f"SELECT {', '.join(columns)} FROM records ORDER BY id")]
    scans = [
        {
            "id": row["id"],
            "source_file": row["source_file"],
            "engine": row["engine"],
            "created_at": row["created_at"],
            "tables": json.loads(row["tables_json"]),
        }
        for row in db.execute("SELECT id, source_file, engine, tables_json, created_at FROM scans ORDER BY id")
    ]

    sources = {r["source_file"] for r in records + scans if r["source_file"]}
    if len(sources) == 1:
        export_stem = Path(next(iter(sources))).stem
    elif sources:
        export_stem = "records_multiple_sources"
    else:
        export_stem = "records"
    download_name = f"{export_stem}.{fmt}"

    n_rows = len(records) + sum(len(t["rows"]) for s in scans for t in s["tables"])
    confirmation = f"Exported {n_rows} row(s) as {download_name}"
    if len(sources) == 1:
        confirmation += f" (source: {next(iter(sources))})"
    flash(confirmation + ".")

    if fmt == "xlsx":
        return Response(
            _excel_workbook(records, columns, scans),
            mimetype="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
            headers={"Content-Disposition": f"attachment; filename={download_name}"},
        )

    if fmt == "csv":
        buf = io.StringIO()
        writer = csv.writer(buf)
        if records or not scans:
            writer.writerow(columns)
            writer.writerows([r[c] for c in columns] for r in records)
        for scan in scans:
            for table in scan["tables"]:
                if buf.tell():
                    writer.writerow([])
                label = scan["source_file"] or f"scan {scan['id']}"
                writer.writerow([f"{label} - {table['title']}" if table["title"] else label])
                writer.writerow(table["columns"])
                writer.writerows(table["rows"])
        return Response(
            buf.getvalue(),
            mimetype="text/csv",
            headers={"Content-Disposition": f"attachment; filename={download_name}"},
        )

    return Response(
        json.dumps({"records": records, "scans": scans}, indent=2, ensure_ascii=False),
        mimetype="application/json",
        headers={"Content-Disposition": f"attachment; filename={download_name}"},
    )


@app.route("/uploads/<path:filename>")
def uploaded_file(filename: str):
    """Serve a raw upload for preview in the /process page."""
    return send_from_directory(str(UPLOADS_DIR), filename)


@app.route("/processed/<path:filename>")
def processed_file(filename: str):
    """Serve a preprocessed image for preview in the /process page."""
    return send_from_directory(str(PROCESSED_DIR), filename)


init_db()


if __name__ == "__main__":
    debug = os.environ.get("FLASK_DEBUG", "0") == "1"
    # 0.0.0.0:7860 matches the Hugging Face Spaces container convention.
    app.run(host="0.0.0.0", port=7860, debug=debug)
